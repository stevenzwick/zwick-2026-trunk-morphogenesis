#!/usr/bin/env python3
"""Compare the bulk RNA-seq values the notebooks plot with the paper's published Source Data.

Run bulkseq/run_bulk_notebooks.py first; it writes the values this script reads.

ED 3b, 3c and 3j are compared value by value against their sheets in the Extended Data Fig.3 workbook.
ED 10g is checked differently, because its sheet publishes CPM while the panel draws a per-gene z-score
of log2(CPM + 1): the sheet's CPM is compared with the notebook's, and then the panel is rebuilt from
the sheet alone and compared with what the notebook plotted. The second half is the one that matters --
it shows a reader can regenerate the heatmap from the published numbers.

    SOURCE_DATA_DIR=/path/to/source_data python bulkseq/verify_bulk_panels.py
"""
from __future__ import annotations
import os, re, sys
from pathlib import Path
import numpy as np
import pandas as pd

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "tools"))
import verify_against_source_data as V  # noqa: E402

BOOK = "Source Data Extended Data Fig.3.xlsx"
PANELS = {"ED 3b": ("Extended Data Fig.3b", "ed3b"),
          "ED 3c": ("Extended Data Fig.3c", "ed3c"),
          "ED 3j": ("Extended Data Fig.3j", "ed3j")}

#: Source Data group label -> GEO sample (library) name. Labels with no rule are not plotted.
GROUP_RULES = [
    (r"^0 hr pSMAD_(\d)$", "pSMAD_0h_rep{}"),
    (r"^7 hr pSMAD_(\d)$", "pSMAD_7h_rep{}"),
    (r"^24 hr pSMAD_(\d)$", "pSMAD_24h_rep{}"),
    # ED 3c's sheet labels its two conditions the way the panel does, by whether the Activin/Nodal
    # inhibitor was present, rather than by the internal library names the GEO samples carry.
    (r"^w/ A83-01_(\d)$", "day2diff_rep{}"),
    (r"^No A83-01_([12])$", "day2_endoderm_rep{}"),
    (r"^BMPR1Akd_(\d)$", "BMPR1Akd_rep{}"),
    (r"^SCRAMBLEkd_(\d)$", "SCRAMBLE_rep{}"),
]

#: Source Data rows each panel deliberately does not plot. Declared rather than inferred: rows
#: with no GROUP_RULES match are dropped before the comparison, so anything outside this list is
#: a row the check would otherwise pass over without ever looking at it.
NOT_PLOTTED = {
    "ED 3b": set(),
    "ED 3c": {"No A83-01_3"},                           # the panel's n is 2; see the lane README
    "ED 3j": {"no KD_1", "no KD_2", "no KD_3"},         # ED 3j is the two-bar KD comparison
}


#: ED 10g's sheet lives in a different workbook and a different shape from ED 3's: a gene x sample
#: matrix with merged group headers over the three condition blocks, rather than one row per value.
ED10G_BOOK = "Source Data Extended Data Fig.10.xlsx"
ED10G_SHEET = "Extended Data Fig.10g"
#: sheet column header -> the notebook's column name for the same library
ED10G_COLUMNS = {"pluripotent_1": "pluripotent_sample_1", "pluripotent_2": "pluripotent_sample_2",
                 "pluripotent_3": "pluripotent_sample_3",
                 "NMP_1": "NMP_1", "NMP_2": "NMP_2", "NMP_3": "NMP_3",
                 "LPM_1": "LPM_1", "LPM_2": "LPM_2", "LPM_3": "LPM_3"}


def read_ed10g_sheet(directory: Path) -> pd.DataFrame:
    """The published matrix, indexed by gene, with the notebook's column names."""
    import openpyxl
    rows = [list(r) for r in openpyxl.load_workbook(directory / ED10G_BOOK, read_only=True)[ED10G_SHEET]
            .iter_rows(values_only=True)]
    head = next(i for i, r in enumerate(rows) if r and r[0] == "gene_name")
    names = [c for c in rows[head] if c is not None]
    body = [r[:len(names)] for r in rows[head + 1:] if r and r[0] is not None]
    table = pd.DataFrame(body, columns=names).set_index("gene_name")
    return table.rename(columns=ED10G_COLUMNS).astype(float)


def row_zscore(frame: pd.DataFrame) -> pd.DataFrame:
    """The notebook's own transform: log2(CPM + 1), then z-score across each gene's samples.

    A gene with no spread would divide by zero, so the notebook leaves that row at 0; matching that
    here keeps the comparison about the data rather than about a convention.
    """
    log2 = np.log2(frame + 1.0)
    std = log2.std(axis=1)
    centred = log2.sub(log2.mean(axis=1), axis=0)
    return centred.div(std.replace(0.0, np.nan), axis=0).fillna(0.0)


def check_ed10g(directory: Path, tol_digits_from: pd.Series) -> tuple[bool, int, float, float, list[str]]:
    """Compare the sheet with the notebook, then rebuild the panel from the sheet alone."""
    published = read_ed10g_sheet(directory)
    cpm = pd.read_csv(REPO / "bulkseq" / "output" / "combined_cpm_expression_matrix_filtered_mean2_TFonly.tsv",
                      sep="\t").set_index("gene_name")
    plotted = pd.read_csv(REPO / "bulkseq" / "output" / "ed10g_heatmap_zscores.tsv", sep="\t", index_col=0)
    plotted.index = plotted.index.astype(str)
    notes: list[str] = []

    missing = [g for g in published.index if g not in cpm.index]
    extra = [g for g in plotted.index if g not in published.index]
    if missing:
        notes.append(f"{len(missing)} sheet genes absent from the notebook: " + ", ".join(missing[:4]))
    if extra:
        notes.append(f"{len(extra)} plotted genes absent from the sheet: " + ", ".join(extra[:4]))
    if list(published.index) != list(plotted.index):
        notes.append("gene order differs from the plotted order")

    genes = [g for g in plotted.index if g in published.index]
    cols = list(plotted.columns)
    tol = V.tolerance_from_precision(tol_digits_from)
    # V.compare works on one dimension, and it pairs values by position, so both sides are flattened
    # through the same gene-by-column order rather than relying on the two frames already agreeing.
    def flat(frame: pd.DataFrame) -> pd.Series:
        return frame.loc[genes, cols].stack()

    cpm_result = V.compare(flat(published), flat(cpm), tol=tol)
    # The sheet publishes CPM, so its rounding is what limits how exactly the panel can be rebuilt.
    # Deriving the tolerance from the z-scores instead would demand precision the sheet never carried.
    z_tol = max(tol, 1e-3)
    z_result = V.compare(flat(row_zscore(published.loc[genes, cols])), flat(plotted), tol=z_tol)
    ok = (not notes and cpm_result["pass"] and z_result["pass"]
          and cpm_result["n"] > 0 and z_result["n"] > 0)
    if not cpm_result["pass"]:
        notes.append("sheet CPM does not match the notebook")
    if not z_result["pass"]:
        notes.append("panel does not rebuild from the sheet")
    return ok, cpm_result["n"], cpm_result["max_abs_diff"], z_result["max_abs_diff"], notes


def sample_for(label: str) -> str | None:
    for pattern, template in GROUP_RULES:
        m = re.match(pattern, label)
        if m:
            return template.format(m.group(1))
    return None


def read_sheet(directory: Path, sheet: str) -> pd.DataFrame:
    import openpyxl
    rows = [list(r) for r in openpyxl.load_workbook(directory / BOOK, read_only=True)[sheet]
            .iter_rows(values_only=True)]
    head = next(i for i, r in enumerate(rows) if r and r[0] == "gene_id")
    body = [r[:4] for r in rows[head + 1:] if r and r[0] is not None]
    return pd.DataFrame(body, columns=["gene_id", "gene_name", "group", "cpm"])


def main() -> int:
    directory = os.environ.get("SOURCE_DATA_DIR")
    if not directory:
        print("Set SOURCE_DATA_DIR to the folder holding the published workbooks.", file=sys.stderr)
        return 2
    outputs = [REPO / "bulkseq" / "output" / f"{stem}_plotted_values.tsv" for _, stem in PANELS.values()]
    if not any(t.is_file() for t in outputs):
        print("No plotted values in bulkseq/output/. Run bulkseq/run_bulk_notebooks.py first.", file=sys.stderr)
        return 2
    print(f"{'panel':7s} {'n':>3s} {'max|diff|':>11s} {'tol':>9s}  result   notes")
    passed = 0
    for panel, (sheet, stem) in PANELS.items():
        table = REPO / "bulkseq" / "output" / f"{stem}_plotted_values.tsv"
        if not table.is_file():
            print(f"{panel:7s} {'':>3s} {'':>11s} {'':>9s}  MISSING  run bulkseq/run_bulk_notebooks.py first")
            continue
        published = read_sheet(Path(directory), sheet)
        published["sample"] = published["group"].astype(str).map(sample_for)
        not_plotted = sorted(published.loc[published["sample"].isna(), "group"].astype(str))
        # A row whose group label matches no rule is dropped here, before the merge, so it can
        # never turn up as an unmatched row. Without an explicit list of the ones the panel is
        # meant to leave out, the check would pass on every row it silently discarded.
        unexpected = sorted(set(not_plotted) - NOT_PLOTTED[panel])
        published = published.dropna(subset=["sample"])
        plotted = pd.read_csv(table, sep="\t")
        merged = plotted.merge(published, on=["gene_name", "sample"], how="outer",
                               suffixes=("_plotted", "_published"), indicator=True)
        unmatched = merged[merged["_merge"] != "both"]
        both = merged[merged["_merge"] == "both"]
        ids_agree = bool((both["gene_id_plotted"].astype(str) == both["gene_id_published"].astype(str)).all())
        tol = V.tolerance_from_precision(published["cpm"].astype(str))
        # Through V.compare so that a value blank on one side is a mismatch, not a skipped row.
        result = V.compare(both["cpm_published"], both["cpm_plotted"], tol=tol)
        ok = (unmatched.empty and not unexpected and ids_agree
              and result["n"] > 0 and result["pass"])
        passed += ok
        notes = []
        if not_plotted:
            notes.append("sheet rows not plotted: " + ", ".join(dict.fromkeys(not_plotted)))
        if unexpected:
            notes.append("UNDECLARED rows not plotted: " + ", ".join(unexpected))
        if not unmatched.empty:
            notes.append(f"{len(unmatched)} unmatched rows")
        if not ids_agree:
            notes.append("gene_id mismatch")
        print(f"{panel:7s} {result['n']:3d} {result['max_abs_diff']:11.3g} {tol:9.2g}  "
              f"{'PASS' if ok else 'FAIL':7s}  {'; '.join(notes)}")

    total = len(PANELS)
    ed10g_values = REPO / "bulkseq" / "output" / "ed10g_heatmap_zscores.tsv"
    if not (Path(directory) / ED10G_BOOK).is_file():
        print(f"{'ED 10g':7s} {'':>3s} {'':>11s} {'':>9s}  MISSING  {ED10G_BOOK} not in SOURCE_DATA_DIR")
    elif not ed10g_values.is_file():
        print(f"{'ED 10g':7s} {'':>3s} {'':>11s} {'':>9s}  MISSING  run bulkseq/run_bulk_notebooks.py first")
    else:
        total += 1
        sheet = read_ed10g_sheet(Path(directory))
        flat = sheet.stack().astype(str)
        ok, n, cpm_diff, z_diff, notes = check_ed10g(Path(directory), flat)
        passed += ok
        notes.append(f"rebuilding the panel from the sheet: max|z diff| {z_diff:.3g}")
        print(f"{'ED 10g':7s} {n:3d} {cpm_diff:11.3g} {V.tolerance_from_precision(flat):9.2g}  "
              f"{'PASS' if ok else 'FAIL':7s}  {'; '.join(notes)}")

    print(f"\n{passed}/{total} bulk panels reproduce")
    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
